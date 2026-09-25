#!/usr/bin/env python3
"""FarmBro robot product-page design brief (review before build)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "FarmBro_Robot_Page_Design_Brief.docx"


def h(doc, text, level=1):
    doc.add_heading(text, level=level)


def p(doc, text, bold=False, italic=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    return para


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, head in enumerate(headers):
        t.rows[0].cells[i].text = head
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            t.rows[r + 1].cells[c].text = cell
    doc.add_paragraph()


def main():
    doc = Document()
    title = doc.add_heading("FarmBro — Robot page design brief", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(
        "Section-by-section brief for /farmbro/:slug. Review this before any page is rebuilt.\n"
        "Status: proposal only. Superdesign and MagicPath were not signed in when this was written. Firecrawl was not connected; competitor notes come from public pages."
    )
    run.italic = True
    run.font.size = Pt(11)

    h(doc, "1. Decision this brief assumes", 1)
    bullets(doc, [
        "Keep the current brand: Terra Jade #1B8F6A, near-black #111311, ivory surfaces, Space Grotesk + IBM Plex Mono.",
        "One shared template for all four robots. Story, numbers, media, and chips change per machine.",
        "Primary action is Enquire / Book a field demo, not a cart.",
        "Do not invent runtime, acres-per-hour, or grower quotes. Unpublished figures stay “On request”.",
        "Build the hybrid mulcher page first, then apply the template to the other three.",
    ])

    h(doc, "2. Why the current page feels simple", 1)
    p(doc, "ProductDetail.tsx is a spec sheet: breadcrumb, name, short body, order button, gallery, four highlight cards, a short spec list, and related machines. The same layout serves a heavy mulcher, a 6×6 carrier, a slope mini, and a drone.")
    p(doc, "A buyer still cannot tell whether the machine fits their crop, what a working day looks like, which tools fit, or how it is safer than a tractor.")

    h(doc, "3. Competitor patterns to borrow", 1)
    table(doc, ["Reference", "Pattern to use"], [
        ["Naïo OZ — naio-technologies.com/en/oz-robot", "Job headline, four benefits, a few big numbers, tool row, short FAQ."],
        ["Carbon Robotics LaserWeeder G2", "Outcome first, then technology modules, then size variants."],
        ["farm-ng Amiga", "Platform story, then a real spec block, then what is included."],
        ["Burro model pages", "One model per URL, units in the spec table, options called out, “who it is for”."],
        ["AgXeed", "Machine plus the system around it: plan, monitor, record."],
    ])
    p(doc, "Use these as structure references only. Do not copy photography, claims, or layout pixel-for-pixel.")

    h(doc, "4. Page sequence", 1)
    p(doc, "Desktop and mobile use the same order. Width changes the layout, not the story.")
    bullets(doc, [
        "00 Sticky enquire bar (mobile; desktop uses the existing navbar Order button)",
        "01 Hero",
        "02 Fit strip",
        "03 Where it works",
        "04 A day on this machine",
        "05 Numbers",
        "06 Compatible attachments",
        "07 Safety versus a tractor",
        "08 Who it is for",
        "09 Gallery",
        "10 Machine FAQ",
        "11 Close + phone / WhatsApp",
        "12 Related machines",
    ])

    h(doc, "5. Section-by-section", 1)

    h(doc, "00 — Sticky enquire bar", 2)
    bullets(doc, [
        "Mobile only, fixed to the bottom, above the WhatsApp badge so they do not overlap.",
        "Left: configuration chip (4X4 UGV, 6X6 UGV, or Scouting drone). Right: Enquire, min height 48px.",
        "Hidden when the hero enquire button is on screen, so the page does not show two primary buttons at once.",
        "Desktop: no extra bar. Navbar already has Order Now.",
    ])

    h(doc, "01 — Hero", 2)
    bullets(doc, [
        "Desktop: two columns. Left is type and actions. Right is the gallery, video first.",
        "Mobile: media on top at about 70vh, then the name and buttons. Do not lock the hero to 100vh.",
        "Eyebrow: tier (Agriculture platform or Crop intelligence drone) with a jade status dot.",
        "H1: machine name. Subline: the existing tagline.",
        "One sentence under that: the existing body, shortened if it runs past four lines on a phone.",
        "Buttons: Enquire (primary, opens the order dialog with this robot id) and Book a field demo (secondary, same dialog with a demo note, or scrolls to the close band).",
        "Under the buttons: price label and availability. Keep “Price on request”.",
        "Gallery: large frame plus thumbnail strip. On mobile the strip is a horizontal scroller. Only the active video plays.",
    ])

    h(doc, "02 — Fit strip", 2)
    p(doc, "Four facts, immediately under the hero. This is the first proof, before the long spec list.")
    table(doc, ["Machine", "Fact 1", "Fact 2", "Fact 3", "Fact 4"], [
        ["Hybrid mulcher", "4X4 UGV", "Hybrid", "Remote controlled", "Heavy mulching"],
        ["6×6 carrier", "6X6 UGV", "Remote controlled", "Mulch · spray · carry", "Estate rows"],
        ["Mini electric", "4X4 UGV", "Electric", "Up to 45°", "Tight ground"],
        ["Canopy Scout", "Scouting drone", "Survey pass", "Health · stress · pest", "Pairs with the ground line"],
    ])
    bullets(doc, [
        "Desktop: one row of four. Mobile: 2×2 grid.",
        "Label in mono, 10px, tracked. Value in the body face, 16–18px.",
        "If a fact is not published, the cell says On request. Do not estimate.",
    ])

    h(doc, "03 — Where it works", 2)
    p(doc, "Chips for terrain and crop. Active chips are jade. Inactive chips are outlined and not implied as a fit.")
    table(doc, ["Machine", "Active chips"], [
        ["Hybrid mulcher", "Open field, Orchard blocks, Long mulching day"],
        ["6×6 carrier", "Plantation rows, Estate tracks, Spray window, Harvest carry"],
        ["Mini electric", "Terraces, Orchard basins, Narrow blocks, Slopes to 45°"],
        ["Canopy Scout", "Before the crew rolls, Canopy health, Water stress, Pest pressure"],
    ])
    p(doc, "Mobile: chips wrap. Do not put them in a horizontal scroll that hides the inactive ones.")

    h(doc, "04 — A day on this machine", 2)
    p(doc, "Three steps. Each step is a still or a short loop from the machine video, plus a caption.")
    bullets(doc, [
        "Hybrid: 1 Set the remote and walk the headland. 2 Mulch the block with the operator off the machine. 3 Park and note what the next pass needs.",
        "Carrier: 1 Choose the job — mulch, spray, or carry. 2 Run the row. 3 Swap the duty without a second vehicle.",
        "Mini: 1 Enter the block a tractor cannot. 2 Hold the slope. 3 Exit without turning the headland into a road.",
        "Scout: 1 Fly the block. 2 Read health, stress, and pest. 3 Send the ground machines only where the map says.",
    ])
    p(doc, "Desktop: three equal columns. Mobile: stacked, image then caption. Do not autoplay three videos.")

    h(doc, "05 — Numbers", 2)
    p(doc, "A dark band, three figures maximum. Only published specs.")
    bullets(doc, [
        "Mini electric may show 45° because that figure is already on the site.",
        "Hybrid may show Hybrid and 4X4. Do not invent litres, hours, or acres.",
        "Carrier payload and tank stay “On request” until the client publishes them.",
        "Scout flight time and camera type stay “On request”.",
        "Each number has a one-line caption so it is not a naked statistic.",
    ])

    h(doc, "06 — Compatible attachments", 2)
    p(doc, "Show only tools that belong on this chassis. Current catalog tools are rotary tiller, boom sprayer, brush cutter, and field trailer.")
    bullets(doc, [
        "Hybrid and mini: tiller, brush cutter. Sprayer and trailer only if the client confirms the rail fits.",
        "Carrier: sprayer and trailer first, then mulching tools.",
        "Scout: no ground implements. Replace this section with “Pairs with” cards linking to the three UGVs.",
        "Each card: code, name, one line, stat, price. Button: Add to enquiry (pre-fills attachment).",
        "Mobile: one card per row. Desktop: up to four.",
    ])

    h(doc, "07 — Safety versus a tractor", 2)
    bullets(doc, [
        "Three points, not a slogan wall.",
        "Operator stays off the machine.",
        "Lighter than a tractor — say zero compaction only if the client will stand behind that line. Otherwise say lighter passes.",
        "Slope: only the mini states 45°. Others say the limit is confirmed on the demo.",
        "Desktop: text left, one still of the remote right. Mobile: still first, then the three points.",
    ])

    h(doc, "08 — Who it is for", 2)
    p(doc, "Three cards. Copy already exists in the dictionaries as audiences, but it is not rendered. Use it here, tuned per machine.")
    bullets(doc, [
        "Farmer — demo on your land. Opens enquiry as a farmer.",
        "Estate or fleet — pilot before a multi-machine order. Opens fleet enquiry.",
        "Government — institutional enquiry. Opens the civil/defence choice, default civil, with defence as a second option in the form.",
        "On the mini page, lead with the farmer card. On the carrier, lead with estate. On the scout, lead with estate (map before the crew).",
    ])

    h(doc, "09 — Gallery", 2)
    bullets(doc, [
        "Keep the bento gallery on desktop.",
        "On mobile, replace bento with a horizontal snap scroller. Bento crops badly under 768px.",
        "First cell is the working video. Later cells are stills.",
        "Alt text stays “{name} — view {n}”.",
    ])

    h(doc, "10 — Machine FAQ", 2)
    p(doc, "Four questions for this robot. Do not repeat the homepage FAQ verbatim.")
    bullets(doc, [
        "Hybrid: How long can it mulch? Who stands where? What fields is it for? What is included in a demo?",
        "Carrier: Can one chassis really spray and carry? What is the tank and payload? How wide are the rows it needs? What does an estate pilot include?",
        "Mini: What slope is published? Where can a tractor not go? Is it quiet enough inside an orchard? How is it charged?",
        "Scout: Is the survey flown by the same crew? What does the map show? Does it replace the mulcher? What licence or permission should we plan for?",
    ])
    p(doc, "Answers that are not confirmed end with: We confirm this on the call.")

    h(doc, "11 — Close", 2)
    bullets(doc, [
        "Dark or jade band, same family as the current order strip.",
        "Heading: Put the {configuration} to work on your rows.",
        "Sub: availability line already in the dictionary.",
        "Actions: Enquire, Call +91 91541 53925, WhatsApp.",
        "Replace hello@farmbro.example before launch. Do not design a new email treatment around the placeholder.",
    ])

    h(doc, "12 — Related machines", 2)
    bullets(doc, [
        "Keep two or three cards.",
        "Order them by the closest alternative: hybrid ↔ mini (both 4×4 mulchers), carrier next to hybrid, scout last as the pair, not a substitute.",
        "Mobile: stacked rows with a 96px thumbnail. Desktop: the current horizontal card is fine.",
    ])

    h(doc, "6. Responsive rules", 1)
    table(doc, ["Rule", "Desktop ≥1024", "Mobile <768"], [
        ["Hero", "Type left, media right", "Media first, type below, ~70vh media"],
        ["Fit strip", "4 columns", "2×2"],
        ["Day steps", "3 columns", "Stacked, one video or still"],
        ["Gallery", "Bento", "Horizontal snap scroll"],
        ["Enquire", "Navbar + hero + close", "Hero + close + sticky bar"],
        ["Type", "Existing clamp sizes", "H1 no smaller than ~32px"],
        ["Tap targets", "Existing buttons", "Minimum 48px height"],
        ["Motion", "Short fades", "No parallax; honor prefers-reduced-motion"],
        ["Video", "Active slide only", "Active slide only; poster if reduced motion"],
    ])
    p(doc, "Check 390px and 768px. The floating dock nav must not cover the hero title or the sticky enquire bar.")

    h(doc, "7. What stays in code", 1)
    bullets(doc, [
        "Route stays /farmbro/:slug. No new URLs for v1 of this template.",
        "Data stays in catalog.ts plus the i18n machines dictionaries. Add optional fields only when a section needs them: terrain chips, day steps, faq, related order.",
        "Enquiry stays on OrderFormDialog via openOrderForm(robot.id).",
        "Do not add Razorpay to this template.",
    ])

    h(doc, "8. Build order after approval", 1)
    bullets(doc, [
        "Approve this brief, or mark sections to cut.",
        "Implement the template on the hybrid mulcher and check it at 390px and 1280px.",
        "Fill the other three robots with their chips, day steps, and FAQs.",
        "Then, if wanted, generate a Superdesign or MagicPath frame of the approved hybrid page. Those tools were not available in the session that wrote this brief.",
    ])

    h(doc, "9. Open questions for the client", 1)
    bullets(doc, [
        "Which attachments officially fit the hybrid, the mini, and the 6×6?",
        "May we publish 45° only on the mini, or do other machines have a slope figure?",
        "Is “no soil compaction” an approved claim?",
        "Are there two real field visits we can name, or do we ship without quotes?",
        "What should Book a field demo create: the same enquiry with a demo flag, or a separate form?",
    ])

    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = end.add_run("Review document. No pages were changed.")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x64, 0x73, 0x6C)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
